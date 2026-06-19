import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy.orm import Session

from app.crud.crud_member import member_crud
from app.models.member import Member


def generate_template() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla Miembros"

    headers = [
        "DNI *",
        "Nombre *",
        "Apellidos",
        "Email",
        "Telefono",
        "Codigo Unico",
        "Fecha Nacimiento (YYYY-MM-DD)",
        "Observaciones Medicas",
        "Otras Observaciones",
    ]

    hf = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    tb = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    ws.merge_cells("A1:I1")
    ws["A1"] = "Plantilla de importacion de miembros"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:I2")
    ws["A2"] = "Los campos con * son obligatorios. Complete una fila por miembro."
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(size=9, italic=True, color="666666")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = tb

    # Example row
    example = ["72722022", "Joaquin", "Rodriguez", "joaquin@mail.com", "999888777", "01", "2000-01-15", "", ""]
    for col, val in enumerate(example, 1):
        cell = ws.cell(row=5, column=col, value=val)
        cell.border = tb
        cell.font = Font(size=10, color="999999", italic=True)

    widths = [15, 20, 20, 28, 15, 15, 28, 30, 30]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def import_members_from_excel(file: io.BytesIO, db: Session) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(file)
    ws = wb.active
    if not ws:
        return {"success": False, "error": "El archivo no contiene hojas"}

    created = 0
    errors = []
    headers_row = 4

    for row_idx, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        dni = str(row[0]).strip() if row[0] else ""
        nombre = str(row[1]).strip() if row[1] else ""

        if not dni or not nombre:
            errors.append(f"Fila {row_idx}: DNI y Nombre son obligatorios")
            continue

        if member_crud.get_by_dni(db, dni):
            errors.append(f"Fila {row_idx}: El DNI '{dni}' ya existe")
            continue

        apellidos = str(row[2]).strip() if row[2] and str(row[2]).strip() != "None" else None
        email = str(row[3]).strip() if row[3] and str(row[3]).strip() != "None" else None
        telefono = str(row[4]).strip() if row[4] and str(row[4]).strip() != "None" else None
        codigo_unico = str(row[5]).strip() if row[5] and str(row[5]).strip() != "None" else None
        fecha_str = str(row[6]).strip() if row[6] and str(row[6]).strip() != "None" else None
        obs_med = str(row[7]).strip() if row[7] and str(row[7]).strip() != "None" else None
        otras_obs = str(row[8]).strip() if row[8] and str(row[8]).strip() != "None" else None

        fecha_nac = None
        if fecha_str:
            try:
                fecha_nac = datetime.strptime(fecha_str, "%Y-%m-%d").date()
            except ValueError:
                errors.append(f"Fila {row_idx}: Fecha '{fecha_str}' no valida (use YYYY-MM-DD)")

        member = Member(
            dni=dni,
            nombre=nombre,
            apellidos=apellidos,
            email=email,
            telefono=telefono,
            codigo_unico=codigo_unico,
            fecha_nacimiento=fecha_nac,
            observaciones_medicas=obs_med,
            otras_observaciones=otras_obs,
            is_active=True,
        )
        db.add(member)
        created += 1

    if created > 0:
        db.commit()

    return {
        "success": True,
        "created": created,
        "errors": errors,
    }
