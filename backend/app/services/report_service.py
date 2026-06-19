import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table as RLTable,
    TableStyle,
)
from reportlab.lib import colors as rl_colors
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.crud.crud_attendance import attendance_crud
from app.crud.crud_member import member_crud
from app.crud.crud_pool_config import pool_config_crud
from app.models.attendance import Attendance
from app.models.member import Member


def get_encabezado(db: Session | None) -> dict:
    if not db:
        return {}
    config = pool_config_crud.get_singleton(db)
    return {
        "nombre": config.institucion_nombre,
        "direccion": config.institucion_direccion,
        "telefono": config.institucion_telefono,
        "email": config.institucion_email,
    }


def get_report_data(
    db: Session,
    fecha_inicio: date | None = None,
    fecha_fin: date | None = None,
    member_id: str | None = None,
):
    if not fecha_inicio:
        fecha_inicio = date.today().replace(day=1)
    if not fecha_fin:
        fecha_fin = date.today()

    query = (
        db.query(
            Member.id,
            Member.dni,
            Member.nombre,
            Member.apellidos,
            Member.telefono,
            func.count(Attendance.id).label("total"),
            func.array_agg(func.date(Attendance.fecha)).label("fechas"),
        )
        .outerjoin(
            Attendance,
            (Attendance.member_id == Member.id)
            & (func.date(Attendance.fecha) >= fecha_inicio)
            & (func.date(Attendance.fecha) <= fecha_fin),
        )
        .filter(Member.is_active == True)
    )

    if member_id:
        query = query.filter(Member.id == member_id)

    query = query.group_by(Member.id).order_by(Member.nombre)
    rows = query.all()

    total_miembros = len(rows)
    miembros_con = sum(1 for r in rows if r.total > 0)
    miembros_sin = total_miembros - miembros_con

    detalle = []
    for r in rows:
        fechas = list(
            sorted(
                set(
                    f.strftime("%Y-%m-%d")
                    for f in (r.fechas or [])
                    if f is not None
                )
            )
        )
        detalle.append(
            {
                "member_id": str(r.id),
                "dni": r.dni,
                "nombre": r.nombre,
                "apellidos": r.apellidos,
                "telefono": r.telefono,
                "total_asistencias": r.total or 0,
                "fechas_asistio": fechas,
            }
        )

    return {
        "summary": {
            "total_miembros_activos": total_miembros,
            "total_asistencias_en_rango": sum(
                r.total or 0 for r in rows
            ),
            "miembros_con_asistencia": miembros_con,
            "miembros_sin_asistencia": miembros_sin,
            "fecha_inicio": fecha_inicio.isoformat(),
            "fecha_fin": fecha_fin.isoformat(),
        },
        "detalle": detalle,
    }


def _add_encabezado_excel(ws, enc: dict, merge_end: str):
    row = 1
    if enc.get("nombre"):
        ws.merge_cells(f"A{row}:{merge_end}{row}")
        ws[f"A{row}"] = enc["nombre"]
        ws[f"A{row}"].font = Font(bold=True, size=16, color="1E40AF")
        ws[f"A{row}"].alignment = Alignment(horizontal="center")
        row += 1
    if enc.get("direccion") or enc.get("telefono") or enc.get("email"):
        parts = [enc.get("direccion") or "", enc.get("telefono") or "", enc.get("email") or ""]
        ws.merge_cells(f"A{row}:{merge_end}{row}")
        ws[f"A{row}"] = " | ".join(p for p in parts if p)
        ws[f"A{row}"].font = Font(size=9, color="6B7280")
        ws[f"A{row}"].alignment = Alignment(horizontal="center")
        row += 1
    return row + 1


def _add_encabezado_pdf(elements, styles, enc: dict):
    if enc.get("nombre"):
        elements.append(Paragraph(enc["nombre"], ParagraphStyle("inst", parent=styles["Title"], fontSize=14, spaceAfter=2, textColor=rl_colors.HexColor("#1E40AF"))))
    parts = [enc.get("direccion") or "", enc.get("telefono") or "", enc.get("email") or ""]
    line = " | ".join(p for p in parts if p)
    if line:
        elements.append(Paragraph(line, ParagraphStyle("instsub", parent=styles["Normal"], fontSize=8, spaceAfter=8, textColor=rl_colors.HexColor("#6B7280"))))


def export_excel(data: dict, encabezado: dict | None = None) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Asistencias"
    enc = encabezado or {}

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    r = _add_encabezado_excel(ws, enc, "F")
    ws.merge_cells(f"A{r}:F{r}")
    ws[f"A{r}"] = "Reporte de Asistencias"
    ws[f"A{r}"].font = Font(bold=True, size=14)
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[r].height = 30
    r += 1

    ws.merge_cells(f"A{r}:F{r}")
    ws[f"A{r}"] = (
        f"Periodo: {data['summary']['fecha_inicio']} - {data['summary']['fecha_fin']}"
    )
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    r += 1

    headers = ["DNI", "Nombre", "Apellidos", "Telefono", "Total Asistencias", "Fechas"]
    hr = r
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for i, row in enumerate(data["detalle"], hr + 1):
        valores = [
            row["dni"],
            row["nombre"],
            row["apellidos"] or "",
            row["telefono"] or "",
            row["total_asistencias"],
            ", ".join(row["fechas_asistio"]),
        ]
        for col, val in enumerate(valores, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border = thin_border
            if col == 6:
                cell.alignment = Alignment(wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_pdf(data: dict, encabezado: dict | None = None) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Title"], fontSize=16, spaceAfter=10
    )
    subtitle_style = ParagraphStyle(
        "CustomSubtitle", parent=styles["Normal"], fontSize=10, spaceAfter=20
    )

    elements = []
    _add_encabezado_pdf(elements, styles, encabezado or {})
    elements.append(Paragraph("Reporte de Asistencias", title_style))
    elements.append(
        Paragraph(
            f"Periodo: {data['summary']['fecha_inicio']} - {data['summary']['fecha_fin']}",
            subtitle_style,
        )
    )

    summary = data["summary"]
    summary_text = (
        f"Miembros activos: {summary['total_miembros_activos']} | "
        f"Asistencias en el periodo: {summary['total_asistencias_en_rango']} | "
        f"Miembros que asistieron: {summary['miembros_con_asistencia']} | "
        f"Miembros que no asistieron: {summary['miembros_sin_asistencia']}"
    )
    elements.append(Paragraph(summary_text, subtitle_style))
    elements.append(Spacer(1, 10 * mm))

    table_data = [
        ["DNI", "Nombre", "Apellidos", "Telefono", "Asistencias", "Fechas"]
    ]
    for row in data["detalle"]:
        table_data.append(
            [
                row["dni"],
                row["nombre"],
                row["apellidos"] or "",
                row["telefono"] or "",
                str(row["total_asistencias"]),
                ", ".join(row["fechas_asistio"]),
            ]
        )

    col_widths = [80, 100, 100, 80, 60, 150]
    table = RLTable(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("BACKGROUND", (0, 1), (-1, -1), rl_colors.HexColor("#F3F4F6")),
                ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.grey),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    return buf


# --- Reporte Detalle de Asistencias ---

DIAS_NOMBRE = ["Lun", "Mar", "Mie", "Jue", "Vie", "Sab", "Dom"]


def get_detalle_asistencias(
    db: Session,
    fecha_inicio: date | None = None,
    fecha_fin: date | None = None,
    member_id: str | None = None,
    is_active: bool | None = None,
):
    if not fecha_inicio:
        fecha_inicio = date.today().replace(day=1)
    if not fecha_fin:
        fecha_fin = date.today()

    from app.models.member import Member
    from app.models.shift_assignment import ShiftAssignment
    from app.models.horario import Horario
    from app.models.attendance import Attendance

    q = db.query(Member)
    if is_active is not None:
        q = q.filter(Member.is_active == is_active)
    if member_id:
        q = q.filter(Member.id == member_id)
    members = q.order_by(Member.nombre).all()

    result = []
    for m in members:
        horarios_raw = (
            db.query(ShiftAssignment, Horario)
            .join(Horario, ShiftAssignment.horario_id == Horario.id)
            .filter(
                ShiftAssignment.member_id == m.id,
                ShiftAssignment.fecha_inicio <= fecha_fin,
                ShiftAssignment.fecha_fin >= fecha_inicio,
            )
            .all()
        )

        hor_map: dict = {}
        for sa, h in horarios_raw:
            key = str(h.id)
            if key not in hor_map:
                hor_map[key] = {
                    "horario_nombre": h.nombre,
                    "hora_inicio": h.hora_inicio.strftime("%H:%M"),
                    "hora_fin": h.hora_fin.strftime("%H:%M"),
                    "dias_semana": [],
                    "fecha_inicio": sa.fecha_inicio.isoformat(),
                    "fecha_fin": sa.fecha_fin.isoformat(),
                }
            d = DIAS_NOMBRE[sa.dia_semana]
            if d not in hor_map[key]["dias_semana"]:
                hor_map[key]["dias_semana"].append(d)

        asistencias = []
        atendances_raw = (
            db.query(Attendance)
            .filter(
                Attendance.member_id == m.id,
                func.date(Attendance.fecha) >= fecha_inicio,
                func.date(Attendance.fecha) <= fecha_fin,
            )
            .order_by(Attendance.fecha.desc())
            .all()
        )
        for a in atendances_raw:
            asistencias.append(
                {
                    "fecha": a.fecha.isoformat() if hasattr(a.fecha, "isoformat") else str(a.fecha),
                    "hora_entrada": a.hora_entrada.strftime("%H:%M"),
                    "observacion": a.observacion,
                }
            )

        result.append(
            {
                "member_id": str(m.id),
                "dni": m.dni,
                "nombre": m.nombre,
                "apellidos": m.apellidos,
                "telefono": m.telefono,
                "email": m.email,
                "is_active": m.is_active,
                "fecha_nacimiento": m.fecha_nacimiento.isoformat() if m.fecha_nacimiento else None,
                "observaciones_medicas": m.observaciones_medicas,
                "total_asistencias_en_rango": len(asistencias),
                "horarios_asignados": list(hor_map.values()),
                "asistencias": asistencias,
            }
        )

    return {
        "miembros": result,
        "total_asistencias": sum(r["total_asistencias_en_rango"] for r in result),
        "fecha_inicio": fecha_inicio.isoformat(),
        "fecha_fin": fecha_fin.isoformat(),
    }


def _set_col_widths(ws, widths: dict[int, int]):
    for col, w in widths.items():
        letter = chr(64 + col) if 1 <= col <= 26 else "A"
        ws.column_dimensions[letter].width = w


def export_detalle_excel(data: dict, encabezado: dict | None = None) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Detalle Asistencias"
    enc = encabezado or {}

    hf = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    tb = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    r = _add_encabezado_excel(ws, enc, "L")
    ws.merge_cells(f"A{r}:L{r}")
    ws[f"A{r}"] = "Reporte Detalle de Asistencias"
    ws[f"A{r}"].font = Font(bold=True, size=14)
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    r += 1
    ws.merge_cells(f"A{r}:L{r}")
    ws[f"A{r}"] = f"Periodo: {data['fecha_inicio']} - {data['fecha_fin']}"
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    r += 1

    hr = r
    headers = ["DNI", "Nombre", "Apellidos", "Telefono", "Estado", "Horario", "Dias", "Fecha", "Ingreso", "Observacion"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=col, value=h)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center")
        cell.border = tb

    row = hr + 1
    for mb in data["miembros"]:
        horario_text = "; ".join(
            f"{h['horario_nombre']} ({h['hora_inicio']}-{h['hora_fin']})" for h in mb["horarios_asignados"]
        )
        dias_text = "; ".join(", ".join(h["dias_semana"]) for h in mb["horarios_asignados"])
        atts = mb["asistencias"] or [{"fecha": "", "hora_entrada": "", "observacion": ""}]
        for a in atts:
            vals = [
                mb["dni"], mb["nombre"], mb["apellidos"] or "", mb["telefono"] or "",
                "Activo" if mb["is_active"] else "Inactivo",
                horario_text, dias_text,
                a["fecha"], a["hora_entrada"], a["observacion"] or "",
            ]
            for col, val in enumerate(vals, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.border = tb
            row += 1
        if not mb["asistencias"]:
            row += 1

    for col, w in enumerate([14, 18, 18, 14, 10, 22, 18, 14, 10, 10, 10, 20], 1):
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_detalle_pdf(data: dict, encabezado: dict | None = None) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10*mm, rightMargin=10*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    ts = ParagraphStyle("T", parent=styles["Title"], fontSize=14, spaceAfter=6)
    ss = ParagraphStyle("S", parent=styles["Normal"], fontSize=9, spaceAfter=12)

    elements = []
    _add_encabezado_pdf(elements, styles, encabezado or {})
    elements.append(Paragraph("Reporte Detalle de Asistencias", ts))
    elements.append(Paragraph(f"Periodo: {data['fecha_inicio']} - {data['fecha_fin']}", ss))

    for mb in data["miembros"]:
        elements.append(Paragraph(
            f"<b>{mb['nombre']} {mb['apellidos'] or ''}</b> - DNI: {mb['dni']} - {'Activo' if mb['is_active'] else 'Inactivo'}",
            ParagraphStyle("m", parent=styles["Normal"], fontSize=8, spaceAfter=4, spaceBefore=10)))
        hl = []
        for h in mb["horarios_asignados"]:
            hl.append(f"{h['horario_nombre']} ({h['hora_inicio']}-{h['hora_fin']}) {', '.join(h['dias_semana'])}")
        if hl:
            elements.append(Paragraph("Horarios: " + "; ".join(hl),
                                      ParagraphStyle("h", parent=styles["Normal"], fontSize=7, spaceAfter=4)))
        if mb["asistencias"]:
            tdata = [["Fecha", "Ingreso", "Observacion"]]
            for a in mb["asistencias"]:
                tdata.append([a["fecha"], a["hora_entrada"], a["observacion"] or "-"])
            tbl = RLTable(tdata, colWidths=[60, 50, 120], repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.3, rl_colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]))
            elements.append(tbl)
        else:
            elements.append(Paragraph("<i>Sin asistencias en el periodo</i>",
                                      ParagraphStyle("n", parent=styles["Normal"], fontSize=7, textColor=rl_colors.grey)))

    doc.build(elements)
    buf.seek(0)
    return buf


# --- Reporte Miembros ---

def get_reporte_miembros(db: Session, is_active: bool | None = None):
    from app.models.member import Member
    q = db.query(Member)
    if is_active is not None:
        q = q.filter(Member.is_active == is_active)
    members = q.order_by(Member.is_active.desc(), Member.nombre).all()
    activos = sum(1 for m in members if m.is_active)
    return {
        "total": len(members),
        "activos": activos,
        "inactivos": len(members) - activos,
        "miembros": [{
            "member_id": str(m.id),
            "dni": m.dni,
            "nombre": m.nombre,
            "apellidos": m.apellidos,
            "email": m.email,
            "telefono": m.telefono,
            "fecha_nacimiento": m.fecha_nacimiento.isoformat() if m.fecha_nacimiento else None,
            "is_active": m.is_active,
            "observaciones_medicas": m.observaciones_medicas,
            "otras_observaciones": m.otras_observaciones,
            "created_at": m.created_at.isoformat() if hasattr(m.created_at, 'isoformat') else str(m.created_at),
        } for m in members],
    }


def export_miembros_excel(data: dict, encabezado: dict | None = None) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Miembros"
    enc = encabezado or {}
    hf = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    tb = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    r = _add_encabezado_excel(ws, enc, "H")
    ws.merge_cells(f"A{r}:H{r}")
    ws[f"A{r}"] = "Reporte de Miembros"
    ws[f"A{r}"].font = Font(bold=True, size=14)
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    r += 1
    ws.merge_cells(f"A{r}:H{r}")
    ws[f"A{r}"] = f"Total: {data['total']} | Activos: {data['activos']} | Inactivos: {data['inactivos']}"
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    r += 1
    hr = r
    headers = ["DNI", "Nombre", "Apellidos", "Email", "Telefono", "Estado", "Fecha Nac.", "Obs. Medicas"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=col, value=h)
        cell.font = hf; cell.fill = hfill; cell.alignment = Alignment(horizontal="center"); cell.border = tb
    for i, m in enumerate(data["miembros"], hr + 1):
        for col, val in enumerate(
            [m["dni"], m["nombre"], m["apellidos"] or "", m["email"] or "", m["telefono"] or "",
             "Activo" if m["is_active"] else "Inactivo", m["fecha_nacimiento"] or "", m["observaciones_medicas"] or ""], 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border = tb
    for col, w in zip([1, 2, 3, 4, 5, 6, 7, 8], [14, 20, 20, 25, 14, 10, 14, 30]):
        ws.column_dimensions[chr(64 + col)].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def export_miembros_pdf(data: dict, encabezado: dict | None = None) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12*mm, rightMargin=12*mm,
                            topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    elements = []
    _add_encabezado_pdf(elements, styles, encabezado or {})
    elements.append(Paragraph("Reporte de Miembros", ParagraphStyle("T", parent=styles["Title"], fontSize=14, spaceAfter=6)))
    elements.append(Paragraph(f"Total: {data['total']} | Activos: {data['activos']} | Inactivos: {data['inactivos']}",
                  ParagraphStyle("S", parent=styles["Normal"], fontSize=9, spaceAfter=12)))
    tdata = [["DNI", "Nombre", "Apellidos", "Email", "Telefono", "Estado", "Fecha Nac."]]
    for m in data["miembros"]:
        tdata.append([m["dni"], m["nombre"], m["apellidos"] or "", m["email"] or "", m["telefono"] or "",
                      "Activo" if m["is_active"] else "Inactivo", m["fecha_nacimiento"] or ""])
    tbl = RLTable(tdata, colWidths=[60, 80, 80, 90, 60, 50, 60], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.3, rl_colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ]))
    elements.append(tbl)
    doc.build(elements)
    buf.seek(0)
    return buf
